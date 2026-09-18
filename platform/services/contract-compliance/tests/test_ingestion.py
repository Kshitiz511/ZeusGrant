"""Ingestion: quotas, ordering, and cleanup.

`IngestionService` had no tests. It is the path every customer upload takes,
it now enforces three plan limits, and it writes to object storage *and* the
database — so a failure here either bills wrongly or leaves orphaned bytes.

The quota tests assert the *order* of the checks as much as their outcome. A
limit that fires after the work it was supposed to prevent is not a limit.
"""

from __future__ import annotations

import io

import docx
import pytest
from zeus_contract_compliance.domain import Document
from zeus_contract_compliance.ingestion import (
    DuplicateDocumentError,
    IngestionService,
    IngestQuota,
    QuotaExceededError,
    UploadTooLargeError,
)

TENANT = "11111111-1111-1111-1111-111111111111"
CONTRACT = "22222222-2222-2222-2222-222222222222"


def docx_bytes(chars: int) -> bytes:
    d = docx.Document()
    d.add_paragraph("x" * chars)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def text_of_pages(pages: int) -> bytes:
    """Plain text sized to estimate exactly ``pages`` under DEC-11."""
    from zeus_contract_compliance.documents import CHARS_PER_PAGE

    return b"x" * ((pages - 1) * CHARS_PER_PAGE + 1)


class FakeStorage:
    def __init__(self) -> None:
        self.objects: dict[str, bytes] = {}
        self.deleted: list[str] = []

    async def put(self, bucket: str, key: str, data: bytes, *, content_type: str) -> None:
        self.objects[key] = data

    async def get(self, bucket: str, key: str) -> bytes:
        return self.objects[key]

    async def delete(self, bucket: str, key: str) -> None:
        self.deleted.append(key)
        self.objects.pop(key, None)


class FakeDocuments:
    """Records every call, so check *ordering* is observable."""

    def __init__(
        self,
        *,
        pages_used: int = 0,
        docs_used: int = 0,
        duplicate: Document | None = None,
        fail_create: bool = False,
    ) -> None:
        self.pages_used = pages_used
        self.docs_used = docs_used
        self.duplicate = duplicate
        self.fail_create = fail_create
        self.calls: list[str] = []
        self.created: dict | None = None

    async def documents_this_month(self, tenant_id: str) -> int:
        self.calls.append("documents_this_month")
        return self.docs_used

    async def pages_this_month(self, tenant_id: str) -> int:
        self.calls.append("pages_this_month")
        return self.pages_used

    async def find_by_checksum(self, tenant_id, contract_id, checksum):
        self.calls.append("find_by_checksum")
        return self.duplicate

    async def create(self, **kw) -> Document:
        self.calls.append("create")
        if self.fail_create:
            raise RuntimeError("insert exploded")
        self.created = kw
        return Document(
            id="33333333-3333-3333-3333-333333333333",
            tenant_id=kw["tenant_id"],
            contract_id=kw["contract_id"],
            filename=kw["filename"],
            content_type=kw["content_type"],
            byte_size=kw["byte_size"],
            checksum=kw["checksum"],
            extracted_chars=kw["extracted_chars"],
            pages=kw["pages"],
            page_basis=kw["page_basis"],
        )


class FakeContracts:
    def __init__(self) -> None:
        self.updates: list[dict] = []

    async def update(self, tenant_id, contract_id, fields) -> None:
        self.updates.append(fields)


def service(documents: FakeDocuments, *, max_bytes: int = 10_000_000):
    storage = FakeStorage()
    contracts = FakeContracts()
    svc = IngestionService(
        storage=storage,
        documents=documents,
        contracts=contracts,
        bucket="evidence",
        max_bytes=max_bytes,
    )
    return svc, storage, contracts


async def ingest(svc, data: bytes, *, filename="c.txt", quota=None, replace_body=False):
    return await svc.ingest(
        tenant_id=TENANT,
        contract_id=CONTRACT,
        filename=filename,
        data=data,
        uploaded_by=None,
        replace_body=replace_body,
        quota=quota,
    )


# --- the page count reaches the database ------------------------------------


@pytest.mark.asyncio
async def test_page_count_is_persisted_with_its_basis():
    """The number cannot be billed against if it is not stored."""
    docs = FakeDocuments()
    svc, _, _ = service(docs)
    result = await ingest(svc, text_of_pages(3))
    assert docs.created is not None
    assert docs.created["pages"] == 3
    assert docs.created["page_basis"] == "estimated"
    assert result.document.pages == 3


@pytest.mark.asyncio
async def test_the_migration_backfill_agrees_with_the_code():
    """0006 backfills `ceil(extracted_chars / 3000)` in SQL.

    That divisor is a copy of CHARS_PER_PAGE. If one changes without the
    other, rows written before the change and after it are billed on
    different scales, and nothing at runtime would ever notice.
    """
    import re
    from pathlib import Path

    from zeus_contract_compliance.documents import CHARS_PER_PAGE

    sql = (
        Path(__file__).resolve().parents[1] / "migrations" / "0006_document_pages.sql"
    ).read_text()
    divisors = re.findall(r"extracted_chars::numeric\s*/\s*(\d+)", sql)
    assert divisors, "backfill expression not found in 0006 -- did it move?"
    assert {int(d) for d in divisors} == {CHARS_PER_PAGE}


# --- unlimited is unlimited (DEC-10 / D22) ----------------------------------


@pytest.mark.asyncio
async def test_no_quota_means_no_limit_queries_at_all():
    """An unlimited plan should not pay for the counting, either."""
    docs = FakeDocuments(pages_used=10_000, docs_used=10_000)
    svc, _, _ = service(docs)
    await ingest(svc, text_of_pages(50))
    assert "pages_this_month" not in docs.calls
    assert "documents_this_month" not in docs.calls


@pytest.mark.asyncio
async def test_an_all_none_quota_is_unlimited():
    """Enterprise plans carry no limit rows, so every field arrives as None."""
    docs = FakeDocuments(pages_used=10_000, docs_used=10_000)
    svc, _, _ = service(docs)
    result = await ingest(svc, text_of_pages(50), quota=IngestQuota())
    assert result.document.pages == 50


# --- pages_per_document -----------------------------------------------------


@pytest.mark.asyncio
async def test_document_over_the_per_document_cap_is_refused():
    docs = FakeDocuments()
    svc, storage, _ = service(docs)
    with pytest.raises(QuotaExceededError) as err:
        await ingest(svc, text_of_pages(30), quota=IngestQuota(pages_per_document=25))
    assert err.value.key == "pages_per_document"
    assert err.value.limit == 25
    assert err.value.requested == 30


@pytest.mark.asyncio
async def test_a_document_exactly_at_the_cap_is_accepted():
    """Off-by-one here is the difference between a sale and a support ticket."""
    docs = FakeDocuments()
    svc, _, _ = service(docs)
    result = await ingest(svc, text_of_pages(25), quota=IngestQuota(pages_per_document=25))
    assert result.document.pages == 25


@pytest.mark.asyncio
async def test_a_refused_document_writes_nothing_anywhere():
    """No stored bytes, no row. A rejected upload must leave no trace and
    must not be billed for storage it was never granted."""
    docs = FakeDocuments()
    svc, storage, contracts = service(docs)
    with pytest.raises(QuotaExceededError):
        await ingest(
            svc,
            text_of_pages(30),
            quota=IngestQuota(pages_per_document=25),
            replace_body=True,
        )
    assert storage.objects == {}
    assert "create" not in docs.calls
    assert contracts.updates == []


# --- pages_per_month --------------------------------------------------------


@pytest.mark.asyncio
async def test_monthly_pages_counts_this_document_before_deciding():
    """95 used + a 10-page file against a cap of 100 is a refusal.

    Checking `used >= limit` instead would accept it and land the tenant at
    105 -- the cap would be advisory.
    """
    docs = FakeDocuments(pages_used=95)
    svc, _, _ = service(docs)
    with pytest.raises(QuotaExceededError) as err:
        await ingest(svc, text_of_pages(10), quota=IngestQuota(pages_per_month=100))
    assert err.value.key == "pages_per_month"
    assert err.value.used == 95
    assert err.value.requested == 10
    assert "5 of 100 left" in str(err.value)


@pytest.mark.asyncio
async def test_a_document_that_exactly_fills_the_month_is_accepted():
    docs = FakeDocuments(pages_used=90)
    svc, _, _ = service(docs)
    result = await ingest(svc, text_of_pages(10), quota=IngestQuota(pages_per_month=100))
    assert result.document.pages == 10


@pytest.mark.asyncio
async def test_an_exhausted_month_reports_zero_remaining_not_a_negative():
    docs = FakeDocuments(pages_used=120)
    svc, _, _ = service(docs)
    with pytest.raises(QuotaExceededError) as err:
        await ingest(svc, text_of_pages(1), quota=IngestQuota(pages_per_month=100))
    assert "0 of 100 left" in str(err.value)


# --- documents_per_month ----------------------------------------------------


@pytest.mark.asyncio
async def test_document_count_is_checked_before_the_file_is_parsed():
    """A tenant already at their cap must not pay for the parse.

    Parsing is the expensive part and it is entirely wasted if the upload was
    never going to be accepted.
    """
    docs = FakeDocuments(docs_used=50)
    svc, _, _ = service(docs)
    with pytest.raises(QuotaExceededError) as err:
        await ingest(
            svc,
            b"%PDF-1.4\nnot actually a valid pdf",
            quota=IngestQuota(documents_per_month=50),
        )
    # A parse would have raised DocumentParseError on these bytes. Getting a
    # quota error instead proves the check ran first.
    assert err.value.key == "documents_per_month"
    assert docs.calls == ["documents_this_month"]


@pytest.mark.asyncio
async def test_the_last_document_of_the_month_is_accepted():
    docs = FakeDocuments(docs_used=49)
    svc, _, _ = service(docs)
    result = await ingest(svc, text_of_pages(1), quota=IngestQuota(documents_per_month=50))
    assert result.document.pages == 1


# --- interaction with the rest of ingestion ---------------------------------


@pytest.mark.asyncio
async def test_size_cap_is_enforced_before_any_quota_query():
    """Rejecting on bytes needs no database round trip."""
    docs = FakeDocuments()
    svc, _, _ = service(docs, max_bytes=10)
    with pytest.raises(UploadTooLargeError):
        await ingest(svc, b"x" * 100, quota=IngestQuota(documents_per_month=1))
    assert docs.calls == []


@pytest.mark.asyncio
async def test_duplicate_detection_still_runs_after_the_quota_checks():
    existing = Document(
        id="44444444-4444-4444-4444-444444444444",
        tenant_id=TENANT,
        contract_id=CONTRACT,
        filename="already.txt",
        content_type="text/plain",
        byte_size=10,
        checksum="abc",
        extracted_chars=10,
    )
    docs = FakeDocuments(duplicate=existing)
    svc, storage, _ = service(docs)
    with pytest.raises(DuplicateDocumentError):
        await ingest(svc, text_of_pages(1), quota=IngestQuota(pages_per_month=1000))
    assert storage.objects == {}


@pytest.mark.asyncio
async def test_a_failed_insert_deletes_the_stored_object():
    """Never leave bytes in storage without a row pointing at them: nothing
    would ever reference them again and nothing would ever clean them up."""
    docs = FakeDocuments(fail_create=True)
    svc, storage, _ = service(docs)
    with pytest.raises(RuntimeError):
        await ingest(svc, text_of_pages(1))
    assert storage.deleted, "orphaned object was not cleaned up"
    assert storage.objects == {}


@pytest.mark.asyncio
async def test_replace_body_sets_the_contract_body_to_the_extracted_text():
    docs = FakeDocuments()
    svc, _, contracts = service(docs)
    await ingest(svc, b"Payment terms are net 30.", replace_body=True)
    assert contracts.updates == [
        {"body": "Payment terms are net 30.", "body_source": "document"}
    ]
